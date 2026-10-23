"""
Módulo para ler e escrever arquivos DOCX (Microsoft Word) - corrigido
Preserva runs/formatação e detecta placeholders quebrados entre runs
"""

import re

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    DOCX_SUPORTADO = True
except ImportError as _e:
    DOCX_SUPORTADO = False
    try:
        from logger import log_erro
        log_erro(f"Falha ao importar python-docx: {_e}")
    except Exception:
        pass


PLACEHOLDER_RE = re.compile(r'\{\{\s*([^}]+?)\s*\}\}')


_REFERENCIAS_CABECALHO = (
    ("headerReference", "default", "header"),
    ("headerReference", "first", "first_page_header"),
    ("headerReference", "even", "even_page_header"),
    ("footerReference", "default", "footer"),
    ("footerReference", "first", "first_page_footer"),
    ("footerReference", "even", "even_page_footer"),
)


def _partes_cabecalho_rodape(section):
    """Cabeçalhos/rodapés realmente referenciados pela seção.

    Só acessa ``section.header``/``footer`` quando existe a referência
    correspondente no ``w:sectPr``. Acessar sem referência faz o python-docx
    criar uma parte vazia (e, em builds congelados sem os templates, registrar
    erro). Assim também entram cabeçalhos de primeira página e de páginas
    pares, que antes eram ignorados.
    """
    sect_pr = getattr(section, "_sectPr", None)
    if sect_pr is None:
        return []

    presentes = set()
    for tag, tipo, attr in _REFERENCIAS_CABECALHO:
        for ref in sect_pr.findall(qn(f"w:{tag}")):
            if (ref.get(qn("w:type")) or "default") == tipo:
                presentes.add(attr)
                break

    partes = []
    for _tag, _tipo, attr in _REFERENCIAS_CABECALHO:
        if attr not in presentes:
            continue
        try:
            partes.append(getattr(section, attr))
        except Exception as e:
            try:
                from logger import log_erro
                log_erro(f"DOCX: nao foi possivel acessar {attr}: {e}")
            except Exception:
                pass
    return partes


def _paragrafos_recursivos(raiz, pai):
    """Percorre todos os w:p do XML, inclusive dentro de formas e controles.

    python-docx só expõe os parágrafos do corpo/tabelas/cabeçalhos; textos
    dentro de caixas de texto (w:txbxContent) e controles de conteúdo
    (w:sdtContent) ficam de fora. Aqui varremos o XML e reaproveitamos a
    mesma lógica de runs/estilo.
    """
    if raiz is None:
        return
    for elemento in raiz.iter(qn("w:p")):
        yield Paragraph(elemento, pai)


def _paragrafos_documento(doc):
    """Parágrafos do corpo (formas, controles, tabelas) e cabeçalhos/rodapés."""
    yield from _paragrafos_recursivos(doc.element.body, doc)

    for section in doc.sections:
        for parte in _partes_cabecalho_rodape(section):
            if parte is not None:
                yield from _paragrafos_recursivos(parte._element, parte)


def _nos_texto(paragraph):
    """Nós ``w:t`` que pertencem a este parágrafo.

    Percorre o XML inteiro (não só ``paragraph.runs``) para alcançar textos
    dentro de hyperlinks, ``w:smartTag``, ``w:ins`` e controles de conteúdo.
    Parágrafos aninhados (caixas de texto) são ignorados aqui e visitados por
    conta própria em ``_paragrafos_recursivos``.
    """
    elemento = paragraph._p
    nos = []
    for no in elemento.iter(qn("w:t")):
        dono = no.getparent()
        while dono is not None and dono.tag != qn("w:p"):
            dono = dono.getparent()
        if dono is elemento:
            nos.append(no)
    return nos


def _definir_texto(no, texto):
    no.text = texto
    if texto and texto != texto.strip():
        no.set(qn("xml:space"), "preserve")


def _substituir_nos(nos, dados):
    """Substitui placeholders entre os nós de texto, preservando os runs."""
    alterou = False
    while True:
        textos = [no.text or "" for no in nos]
        completo = "".join(textos)

        alvo = None
        for match in PLACEHOLDER_RE.finditer(completo):
            chave = match.group(1).strip()
            if chave in dados:
                alvo = (match, str(dados[chave]))
                break
        if alvo is None:
            return alterou

        match, valor = alvo
        inicio, fim = match.span()

        faixas = []
        pos = 0
        for indice, texto in enumerate(textos):
            faixas.append((indice, pos, pos + len(texto)))
            pos += len(texto)
        afetados = [f for f in faixas if f[1] < fim and f[2] > inicio]
        if not afetados:
            return alterou
        primeiro = afetados[0][0]

        for indice, ini_no, _fim_no in reversed(afetados):
            texto = textos[indice]
            ini = max(inicio, ini_no) - ini_no
            fim_local = min(fim, ini_no + len(texto)) - ini_no
            novo = (
                texto[:ini] + valor + texto[fim_local:]
                if indice == primeiro
                else texto[:ini] + texto[fim_local:]
            )
            _definir_texto(nos[indice], novo)

        if "".join(no.text or "" for no in nos) == completo:
            return alterou
        alterou = True


def _text_from_paragraph(paragraph):
    """Concatena todo o texto do parágrafo, inclusive hyperlinks."""
    return "".join(no.text or "" for no in _nos_texto(paragraph))


def _replace_in_paragraph(paragraph, dados):
    """Substitui placeholders no nível dos nós de texto.

    Diferente da versão anterior, não concentra o parágrafo em um único run:
    cada valor herda a formatação do trecho onde o placeholder começa e
    hyperlinks/campos continuam existindo.
    """
    nos = _nos_texto(paragraph)
    if not nos:
        return False
    if "{{" not in "".join(no.text or "" for no in nos):
        return False
    return _substituir_nos(nos, dados)


def extrair_placeholders_docx(caminho_docx):
    if not DOCX_SUPORTADO:
        raise Exception("Suporte a DOCX não disponível. Instale: pip install python-docx")

    placeholders = set()
    doc = Document(caminho_docx)

    for paragraph in _paragrafos_documento(doc):
        text = _text_from_paragraph(paragraph)
        for match in PLACEHOLDER_RE.findall(text):
            nome = match.strip()
            if 0 < len(nome) <= 80:
                placeholders.add(nome)

    return placeholders


def gerar_docx_preenchido(caminho_modelo, dados, caminho_saida):
    if not DOCX_SUPORTADO:
        raise Exception("Suporte a DOCX não disponível. Instale: pip install python-docx")

    doc = Document(caminho_modelo)

    for paragraph in _paragrafos_documento(doc):
        _replace_in_paragraph(paragraph, dados)

    doc.save(caminho_saida)


def docx_suportado():
    return DOCX_SUPORTADO
