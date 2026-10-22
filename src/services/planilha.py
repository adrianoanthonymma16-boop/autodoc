"""
PlanilhaService - lê .xlsx/.csv/.tsv e mapeia colunas para placeholders.

Contrato: a primeira linha não vazia é o cabeçalho; cada linha seguinte
vira um documento (uma linha = um documento gerado).
"""
import csv
import os
import re
import unicodedata
from difflib import SequenceMatcher

EXTENSOES = (".xlsx", ".xlsm", ".csv", ".tsv")
LIMIAR_SUGESTAO = 0.75


def normalizar(texto: str) -> str:
    """minúsculo, sem acentos, só [a-z0-9] separados por espaço."""
    texto = unicodedata.normalize("NFKD", str(texto or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", texto.lower()).strip()


def _limpar(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _ler_xlsx(caminho: str) -> list[list[str]]:
    try:
        import openpyxl
    except ImportError as e:
        raise ValueError("Para ler planilhas .xlsx instale: pip install openpyxl") from e

    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        ws = wb.active
        linhas = [[_limpar(c) for c in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return linhas


def _ler_csv(caminho: str) -> list[list[str]]:
    ultimo_erro = None
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with open(caminho, encoding=encoding, newline="") as f:
                amostra = f.read(8192)
                f.seek(0)
                delimitador = ","
                try:
                    delimitador = csv.Sniffer().sniff(amostra, delimiters=",;\t").delimiter
                except csv.Error:
                    if caminho.lower().endswith(".tsv"):
                        delimitador = "\t"
                leitor = csv.reader(f, delimiter=delimitador)
                return [[_limpar(c) for c in linha] for linha in leitor]
        except (UnicodeDecodeError, OSError) as e:
            ultimo_erro = e
            continue
    raise ValueError(f"Não foi possível ler o arquivo: {ultimo_erro}")


def ler_planilha(caminho: str) -> tuple[list[str], list[list[str]]]:
    """Retorna (colunas, linhas_de_dados). Ignora linhas totalmente vazias."""
    ext = os.path.splitext(caminho)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        brutas = _ler_xlsx(caminho)
    elif ext in (".csv", ".tsv"):
        brutas = _ler_csv(caminho)
    else:
        raise ValueError(f"Formato não suportado: {ext}. Use .xlsx, .csv ou .tsv.")

    linhas = [linha for linha in brutas if any(v for v in linha)]
    if not linhas:
        raise ValueError("A planilha está vazia.")

    cabecalho = list(linhas[0])
    while cabecalho and not cabecalho[-1]:
        cabecalho.pop()
    if not cabecalho:
        raise ValueError("A primeira linha (cabeçalho) está vazia.")

    colunas = [c or f"Coluna {i + 1}" for i, c in enumerate(cabecalho)]
    dados = []
    largura = len(colunas)
    for linha in linhas[1:]:
        linha = list(linha[:largura])
        if len(linha) < largura:
            linha += [""] * (largura - len(linha))
        dados.append(linha)
    if not dados:
        raise ValueError("A planilha não tem nenhuma linha de dados.")
    return colunas, dados


def sugerir_mapeamento(colunas: list[str], placeholders: list[str]) -> dict[str, str]:
    """Sugere coluna para cada placeholder (exato, contém, ou similaridade)."""
    sugestoes: dict[str, str] = {}
    usadas: set[str] = set()
    for ph in placeholders:
        alvo = normalizar(ph)
        if not alvo:
            continue
        melhor = None
        melhor_score = 0.0
        for coluna in colunas:
            if coluna in usadas:
                continue
            nome = normalizar(coluna)
            if not nome:
                continue
            if nome == alvo:
                score = 1.0
            elif nome in alvo or alvo in nome:
                score = 0.9
            else:
                score = SequenceMatcher(None, nome, alvo).ratio()
            if score > melhor_score:
                melhor, melhor_score = coluna, score
        if melhor is not None and melhor_score >= LIMIAR_SUGESTAO:
            sugestoes[ph] = melhor
            usadas.add(melhor)
    return sugestoes


def dados_da_linha(colunas: list[str], linha: list[str], mapeamento: dict[str, str]) -> dict[str, str]:
    """Converte uma linha em {placeholder: valor} usando o mapeamento."""
    indice = {nome: i for i, nome in enumerate(colunas)}
    dados = {}
    for ph, coluna in (mapeamento or {}).items():
        i = indice.get(coluna)
        dados[ph] = linha[i] if i is not None and i < len(linha) else ""
    return dados
