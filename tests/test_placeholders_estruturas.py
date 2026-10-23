"""Placeholders em hyperlinks, controle de alterações e cabeçalhos ODT/DOCX."""
import zipfile

from docx import Document
from docx.oxml import parse_xml

from modelo_docx import extrair_placeholders_docx, gerar_docx_preenchido
from modelo_odt import extrair_placeholders_odt, gerar_odt_preenchido

NS_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
NS_R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'


def _docx_com_hyperlink(caminho):
    doc = Document()
    par = doc.add_paragraph()
    par.add_run("Antes ")
    par._p.append(parse_xml(
        f'<w:hyperlink {NS_W} {NS_R} r:id="rId7">'
        '<w:r><w:t>{{link}}</w:t></w:r>'
        '</w:hyperlink>'
    ))
    par.add_run(" depois")
    doc.save(caminho)


def _docx_com_alteracao(caminho):
    doc = Document()
    doc.add_paragraph("Texto ")
    doc.element.body.insert(1, parse_xml(
        f'<w:p {NS_W}><w:ins w:id="1" w:author="revisor" w:date="2026-01-01T00:00:00Z">'
        '<w:r><w:t>{{alterado}}</w:t></w:r>'
        '</w:ins></w:p>'
    ))
    doc.save(caminho)


def test_extrai_e_preenche_placeholder_em_hyperlink(tmp_path):
    entrada = tmp_path / "link.docx"
    saida = tmp_path / "link-preenchido.docx"
    _docx_com_hyperlink(entrada)

    assert extrair_placeholders_docx(entrada) == {"link"}

    gerar_docx_preenchido(entrada, {"link": "https://exemplo.test"}, saida)
    doc = Document(saida)
    corpo = doc.element.body
    assert doc.paragraphs[0].text == "Antes https://exemplo.test depois"
    assert len(corpo.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hyperlink")) == 1


def test_extrai_e_preenche_placeholder_em_alteracao_controlada(tmp_path):
    entrada = tmp_path / "revisao.docx"
    saida = tmp_path / "revisao-preenchida.docx"
    _docx_com_alteracao(entrada)

    assert extrair_placeholders_docx(entrada) == {"alterado"}

    gerar_docx_preenchido(entrada, {"alterado": "novo"}, saida)
    texto = "".join(doc_texto.text or "" for doc_texto in Document(saida).element.body.iter(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
    ))
    assert "novo" in texto


def test_valor_com_barra_invertida_nao_quebra(tmp_path):
    entrada = tmp_path / "barra.docx"
    saida = tmp_path / "barra-preenchida.docx"
    doc = Document()
    doc.add_paragraph("{{caminho}}")
    doc.save(entrada)

    gerar_docx_preenchido(entrada, {"caminho": r"C:\temp\g"}, saida)

    assert Document(saida).paragraphs[0].text == r"C:\temp\g"


def test_docx_cabecalhos_primeira_e_par(tmp_path):
    entrada = tmp_path / "hdr.docx"
    saida = tmp_path / "hdr-preenchido.docx"
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = True
    secao = doc.sections[0]
    secao.different_first_page_header_footer = True
    secao.header.paragraphs[0].text = "{{normal}}"
    secao.first_page_header.paragraphs[0].text = "{{primeira}}"
    secao.even_page_header.paragraphs[0].text = "{{par}}"
    secao.footer.paragraphs[0].text = "{{rodape}}"
    doc.add_paragraph("{{corpo}}")
    doc.save(entrada)

    assert extrair_placeholders_docx(entrada) == {
        "normal", "primeira", "par", "rodape", "corpo"
    }

    gerar_docx_preenchido(
        entrada,
        {"normal": "N", "primeira": "P", "par": "E", "rodape": "R", "corpo": "C"},
        saida,
    )
    novo = Document(saida)
    secao = novo.sections[0]
    assert secao.header.paragraphs[0].text == "N"
    assert secao.first_page_header.paragraphs[0].text == "P"
    assert secao.even_page_header.paragraphs[0].text == "E"
    assert secao.footer.paragraphs[0].text == "R"
    assert novo.paragraphs[0].text == "C"


_ODT_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
 xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 office:version="1.2"><office:body><office:text>
 <text:p>{{corpo}}</text:p>
</office:text></office:body></office:document-content>"""

_ODT_STYLES = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-styles
 xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 office:version="1.2"><office:master-styles>
 <style:master-page style:name="Standard">
  <style:header><text:p>{{cabecalho}}</text:p></style:header>
  <style:footer><text:p>{{rodape}}</text:p></style:footer>
 </style:master-page>
</office:master-styles></office:document-styles>"""


def _odt_com_cabecalho(caminho):
    with zipfile.ZipFile(caminho, "w") as odt:
        odt.writestr(
            "mimetype", "application/vnd.oasis.opendocument.text",
            compress_type=zipfile.ZIP_STORED,
        )
        odt.writestr("content.xml", _ODT_CONTENT)
        odt.writestr("styles.xml", _ODT_STYLES)


def test_odt_extrai_placeholders_de_cabecalho_e_rodape(tmp_path):
    caminho = tmp_path / "hdr.odt"
    _odt_com_cabecalho(caminho)

    assert extrair_placeholders_odt(caminho) == {"corpo", "cabecalho", "rodape"}


def test_odt_preenche_placeholders_de_cabecalho_e_rodape(tmp_path):
    entrada = tmp_path / "hdr.odt"
    saida = tmp_path / "hdr-preenchido.odt"
    _odt_com_cabecalho(entrada)

    gerar_odt_preenchido(
        entrada,
        {"corpo": "C", "cabecalho": "H", "rodape": "R"},
        saida,
    )

    with zipfile.ZipFile(saida) as odt:
        assert odt.namelist()[0] == "mimetype"
        assert ">C<" in odt.read("content.xml").decode("utf-8")
        estilos = odt.read("styles.xml").decode("utf-8")
        assert ">H<" in estilos
        assert ">R<" in estilos
        assert "{{" not in estilos
