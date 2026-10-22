"""Placeholders dentro de formas (caixas de texto) e controles de conteudo."""
import zipfile

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from modelo_docx import extrair_placeholders_docx, gerar_docx_preenchido
from modelo_odt import extrair_placeholders_odt, gerar_odt_preenchido

NS_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
NS_V = 'xmlns:v="urn:schemas-microsoft-com:vml"'


def _paragrafo_com_caixa(texto_ancora, partes):
    runs = "".join(f"<w:r><w:t>{p}</w:t></w:r>" for p in partes)
    return parse_xml(
        f'<w:p {NS_W} {NS_V}>'
        f'<w:r><w:t xml:space="preserve">{texto_ancora}</w:t></w:r>'
        '<w:r><w:pict>'
        '<v:shape id="Caixa1" style="width:120pt;height:60pt">'
        '<v:textbox><w:txbxContent>'
        f'<w:p>{runs}</w:p>'
        '</w:txbxContent></v:textbox>'
        '</v:shape>'
        '</w:pict></w:r>'
        '</w:p>'
    )


def _docx_com_forma(caminho, ancora="{{fora}}", partes=("{{no", "me}}")):
    doc = Document()
    doc.add_paragraph(ancora)
    doc.element.body.append(_paragrafo_com_caixa("Antes ", partes))
    doc.save(caminho)


def _texto_das_formas(doc):
    return "".join(
        t.text or ""
        for caixa in doc.element.body.iter(qn("w:txbxContent"))
        for t in caixa.iter(qn("w:t"))
    )


def test_extrai_placeholder_dentro_de_caixa_de_texto(tmp_path):
    caminho = tmp_path / "forma.docx"
    _docx_com_forma(caminho, ancora="{{fora}}")

    assert extrair_placeholders_docx(caminho) == {"fora", "nome"}


def test_preenche_placeholder_dentro_de_caixa_de_texto(tmp_path):
    entrada = tmp_path / "forma.docx"
    saida = tmp_path / "preenchido.docx"
    _docx_com_forma(entrada)

    gerar_docx_preenchido(entrada, {"fora": "João", "nome": "Maria"}, saida)

    doc = Document(saida)
    assert doc.paragraphs[0].text == "João"
    assert _texto_das_formas(doc) == "Maria"


def test_ancora_com_forma_preserva_a_caixa(tmp_path):
    entrada = tmp_path / "ancora.docx"
    saida = tmp_path / "ancora-preenchida.docx"
    _docx_com_forma(entrada, ancora="Contrato de {{fora}}")

    gerar_docx_preenchido(entrada, {"fora": "Aluguel", "nome": "Ana"}, saida)

    doc = Document(saida)
    assert doc.paragraphs[0].text == "Contrato de Aluguel"
    assert _texto_das_formas(doc) == "Ana"


def test_extrai_e_preenche_controle_de_conteudo(tmp_path):
    entrada = tmp_path / "controle.docx"
    saida = tmp_path / "controle-preenchido.docx"
    doc = Document()
    doc.element.body.append(parse_xml(
        f'<w:sdt {NS_W}><w:sdtContent>'
        '<w:p><w:r><w:t>{{controle}}</w:t></w:r></w:p>'
        '</w:sdtContent></w:sdt>'
    ))
    doc.save(entrada)

    assert extrair_placeholders_docx(entrada) == {"controle"}

    gerar_docx_preenchido(entrada, {"controle": "OK"}, saida)
    texto = "".join(t.text or "" for t in Document(saida).element.body.iter(qn("w:t")))
    assert texto == "OK"


ODT_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
 xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
 xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
 office:version="1.2">
 <office:body><office:text>
  <text:p>Olá {{fora}}</text:p>
  <text:p><draw:frame draw:name="Caixa" svg:width="5cm" svg:height="2cm">
   <draw:text-box><text:p>{{dentro}}</text:p></draw:text-box>
  </draw:frame></text:p>
 </office:text></office:body>
</office:document-content>
"""


def _odt_escrever(caminho, content):
    with zipfile.ZipFile(caminho, "w") as odt:
        odt.writestr(
            "mimetype", "application/vnd.oasis.opendocument.text",
            compress_type=zipfile.ZIP_STORED,
        )
        odt.writestr("content.xml", content, compress_type=zipfile.ZIP_DEFLATED)


def _odt_com_forma(caminho):
    _odt_escrever(caminho, ODT_CONTENT)


def test_extrai_placeholder_dentro_de_forma_odt(tmp_path):
    caminho = tmp_path / "forma.odt"
    _odt_com_forma(caminho)

    assert extrair_placeholders_odt(caminho) == {"fora", "dentro"}


def test_preenche_placeholder_dentro_de_forma_odt(tmp_path):
    entrada = tmp_path / "forma.odt"
    saida = tmp_path / "preenchido.odt"
    _odt_com_forma(entrada)

    gerar_odt_preenchido(entrada, {"fora": "A", "dentro": "B"}, saida)

    with zipfile.ZipFile(saida) as odt:
        content = odt.read("content.xml").decode("utf-8")
    assert "A" in content and "B" in content
    assert "{{fora}}" not in content and "{{dentro}}" not in content


ODT_SPANS = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
 xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
 xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
 office:version="1.2">
 <office:body><office:text>
  <text:p><draw:frame draw:name="Caixa" svg:width="5cm" svg:height="2cm">
   <draw:text-box><text:p>
    <text:span>Prezado {{no</text:span><text:span>me}}</text:span>
   </text:p></draw:text-box>
  </draw:frame></text:p>
  <text:p>Texto <text:span>fixo</text:span> aqui</text:p>
 </office:text></office:body>
</office:document-content>
"""


def test_extrai_placeholder_quebrado_entre_spans_em_forma_odt(tmp_path):
    caminho = tmp_path / "spans.odt"
    _odt_escrever(caminho, ODT_SPANS)

    assert extrair_placeholders_odt(caminho) == {"nome"}


def test_preenche_placeholder_quebrado_entre_spans_em_forma_odt(tmp_path):
    entrada = tmp_path / "spans.odt"
    saida = tmp_path / "spans-preenchido.odt"
    _odt_escrever(entrada, ODT_SPANS)

    gerar_odt_preenchido(entrada, {"nome": "Maria"}, saida)

    with zipfile.ZipFile(saida) as odt:
        content = odt.read("content.xml").decode("utf-8")
    assert "Prezado Maria" in content
    assert "{{no" not in content and "me}}" not in content
    assert "Texto <text:span>fixo</text:span> aqui" in content


def test_varios_placeholders_no_mesmo_no_odt(tmp_path):
    entrada = tmp_path / "varios.odt"
    saida = tmp_path / "varios-preenchido.odt"
    _odt_escrever(
        entrada,
        ODT_CONTENT.replace("<text:p>Olá {{fora}}</text:p>",
                            "<text:p>{{fora}} e {{dentro}}</text:p>"),
    )

    gerar_odt_preenchido(entrada, {"fora": "A", "dentro": "B"}, saida)

    with zipfile.ZipFile(saida) as odt:
        content = odt.read("content.xml").decode("utf-8")
    assert "A e B" in content
    assert "{{fora}}" not in content and "{{dentro}}" not in content
