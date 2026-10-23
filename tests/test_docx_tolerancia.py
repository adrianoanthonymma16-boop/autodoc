"""Cabeçalhos/rodapés do DOCX: referências, primeira página e tolerância."""
from docx import Document
from docx.oxml import parse_xml

from modelo_docx import _partes_cabecalho_rodape

NS_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
NS_R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'


def _secao_falsa_com_referencia(tipo="default"):
    class _Secao:
        @property
        def header(self):
            raise FileNotFoundError("default-header.xml")

    _Secao._sectPr = parse_xml(
        f'<w:sectPr {NS_W} {NS_R}><w:headerReference w:type="{tipo}" '
        'r:id="rId99"/></w:sectPr>'
    )
    return _Secao()


def test_secao_sem_referencia_nao_cria_partes():
    doc = Document()
    secao = doc.sections[0]
    antes = len(doc.part.package.parts)

    assert _partes_cabecalho_rodape(secao) == []
    assert len(doc.part.package.parts) == antes


def test_falha_em_parte_referenciada_nao_propaga():
    assert _partes_cabecalho_rodape(_secao_falsa_com_referencia()) == []


def test_le_cabecalho_e_rodape_default():
    doc = Document()
    secao = doc.sections[0]
    secao.header.paragraphs[0].text = "{{cabecalho}}"
    secao.footer.paragraphs[0].text = "{{rodape}}"

    partes = _partes_cabecalho_rodape(secao)

    assert [p.paragraphs[0].text for p in partes] == ["{{cabecalho}}", "{{rodape}}"]


def test_le_cabecalho_de_primeira_pagina_e_par():
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = True
    secao = doc.sections[0]
    secao.different_first_page_header_footer = True
    secao.header.paragraphs[0].text = "{{normal}}"
    secao.first_page_header.paragraphs[0].text = "{{primeira}}"
    secao.even_page_header.paragraphs[0].text = "{{par}}"

    textos = [p.paragraphs[0].text for p in _partes_cabecalho_rodape(secao)]

    assert textos == ["{{normal}}", "{{primeira}}", "{{par}}"]
